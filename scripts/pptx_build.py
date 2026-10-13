#!/usr/bin/env python3
"""Build a deck from the user's template by filling layout placeholders by idx.

Usage:
  python pptx_build.py --list-layouts TEMPLATE [--json]
  python pptx_build.py TEMPLATE SPEC.json OUT.pptx [--keep-template-slides] [--remove-empty]
                       [--force] [--json]

Needs python-pptx (pip install python-pptx); the other scripts do not.

TEMPLATE: the user's .pptx or .potx (a .potx is read through a copy with the presentation
content type, because python-pptx 1.0.2 refuses the template type: tested). Macro-enabled
templates are refused; make a macro-free copy with pptx_sanitize.py --remove-macros first.

SPEC.json:
  {"slides": [
     {"layout": "Title Slide", "fill": {"0": "Quarterly review", "1": "October 2026"}},
     {"layout": 1, "fill": {"0": "Results", "1": ["Revenue up 12%", {"text": "EMEA led", "level": 1}]},
      "notes": "Mention the one-off in July.", "hidden": false},
     {"layout": "Picture with Caption", "fill": {"0": "Site", "1": {"image": "site.png"}, "2": "Caption"}}
  ]}
  layout: name (case-insensitive) or 0-based index from --list-layouts.
  fill keys: placeholder idx from --list-layouts. Values: a string, a list of paragraphs
  (strings or {"text", "level"}), {"image": path} for a picture placeholder (path relative
  to the spec file), or {"table": [["Region", "Q3"], ["EMEA", "12%"]]} (rows of strings, the
  first row is the header; up to 50 rows and 20 columns). A table placeholder gets the table
  itself; any other placeholder is replaced by a table at its position and size.
Why idx: placeholders keep the layout's position, fonts and bullets, and titles stay titles in
the outline and for screen readers. Text boxes at fixed coordinates do not.

Template slides are removed completely (as pptx_slides.py --delete does), with the template's
sections, unless --keep-template-slides. Placeholders left empty show "Click to add ..." while editing (not in
the show); --remove-empty deletes them. Spec text with control, bidi or zero-width characters
is refused.
Image paths must be relative and stay inside the spec's folder (realpath checked), so a spec
cannot pull an arbitrary file into the deck; --allow-image-outside lifts this after asking.
OUT may not be the template, the spec or a symlink, even with --force. Errors from python-pptx
or the image library come back as exit 2 with a message, not a traceback.

Exit codes: 0 written, 1 spec does not fit the template (unknown layout or placeholder idx,
nothing written), 2 python-pptx missing, bad spec, unsupported template, output exists,
3 safety refusal from the package reader.
"""
import argparse
import json
import os
import shutil
import sys
import tempfile
import zipfile

import pptpkg as pk
import pptx_sanitize as ps


class SpecError(Exception):
    """Spec does not match the template (exit 1)."""


def load_template(path, work):
    pkg = pk.Package(path)  # caps, encrypted, legacy
    if pkg.kind in ("pptm", "ppsm", "potm"):
        raise pk.Unsupported("template is macro-enabled (%s); make a macro-free copy with "
                             "pptx_sanitize.py --remove-macros first" % pkg.kind)
    copy = os.path.join(work, "template.pptx")
    if pkg.kind == "pptx":
        shutil.copyfile(path, copy)
        return copy
    want = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    ct = pkg.ct  # parsed [Content_Types].xml; change only the Override for the main part
    for o in ct.findall("ct:Override", pk.NS):
        if o.get("PartName", "").lstrip("/") == pkg.main:
            o.set("ContentType", want)
    data_ct = pk.etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(copy, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = data_ct if info.filename == "[Content_Types].xml" else zin.read(info.filename)
            zout.writestr(info, data)
    return copy


def layouts_of(prs):
    out = []
    for m_i, master in enumerate(prs.slide_masters):
        for layout in master.slide_layouts:
            phs = [{"idx": ph.placeholder_format.idx, "type": str(ph.placeholder_format.type).split(".")[-1].split(" ")[0],
                    "name": pk.clean_text(ph.name, 80)} for ph in layout.placeholders]
            out.append({"index": len(out), "master": m_i, "name": pk.clean_text(layout.name, 120),
                        "placeholders": phs, "_obj": layout})
    return out


def check(value, where):
    if isinstance(value, str) and pk._UNSAFE.search(value):
        raise pk.Unsupported("%s contains control, bidi or zero-width characters" % where)


def fill_text(tf, value, where):
    paras = value if isinstance(value, list) else [value]
    tf.clear()  # leaves one empty paragraph that keeps the placeholder's formatting
    for i, item in enumerate(paras):
        if isinstance(item, str):
            text, level = item, 0
        elif isinstance(item, dict):
            text, level = item.get("text", ""), item.get("level", 0)
            if not isinstance(text, str) or not isinstance(level, int) or isinstance(level, bool):
                raise pk.Unsupported("%s: text must be a string and level an integer 0-8" % where)
        else:
            raise pk.Unsupported("%s: a paragraph must be a string or {\"text\", \"level\"}" % where)
        check(text, where)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = max(0, min(level, 8))


def fill_table(slide, ph, rows, where):
    """Put a table where placeholder ph is. Rows are lists of strings; row 0 is the header."""
    if (not isinstance(rows, list) or not rows or len(rows) > 50
            or not all(isinstance(r, list) and r for r in rows)):
        raise pk.Unsupported("%s: table must be 1 to 50 rows, each a non-empty list" % where)
    cols = max(len(r) for r in rows)
    if cols > 20:
        raise pk.Unsupported("%s: table has more than 20 columns" % where)
    for r in rows:
        for c in r:
            if not isinstance(c, str):
                raise pk.Unsupported("%s: table cells must be strings" % where)
            check(c, where + " table cell")
    if hasattr(ph, "insert_table"):
        table = ph.insert_table(len(rows), cols).table
    else:
        x, y, w, h = ph.left, ph.top, ph.width, ph.height
        if None in (x, y, w, h):
            raise pk.Unsupported("%s: placeholder has no position to place a table" % where)
        el = ph._element
        el.getparent().remove(el)
        table = slide.shapes.add_table(len(rows), cols, x, y, w, h).table
    for i, r in enumerate(rows):
        for j in range(cols):
            table.cell(i, j).text = r[j] if j < len(r) else ""


def image_path(base, rel, allow_outside, n):
    """Images must sit in the spec's folder or below it, so a spec cannot pull an arbitrary
    file from disk into a deck that will be shared. --allow-image-outside lifts this."""
    if not isinstance(rel, str) or not rel:
        raise pk.Unsupported("slide %d: image must be a path string" % n)
    if not allow_outside and (os.path.isabs(rel) or rel.startswith("~")):
        raise pk.Unsupported("slide %d: image path must be relative to the spec folder: %s "
                             "(--allow-image-outside to permit)" % (n, rel))
    path = os.path.realpath(os.path.join(base, rel))
    root = os.path.realpath(base)
    if not allow_outside and os.path.commonpath([path, root]) != root:
        raise pk.Unsupported("slide %d: image is outside the spec folder: %s "
                             "(--allow-image-outside to permit)" % (n, rel))
    if not os.path.isfile(path):
        raise pk.Unsupported("slide %d: image not found: %s" % (n, rel))
    return path


def build(a):
    try:
        from pptx import Presentation
    except ImportError:
        raise pk.Unsupported("python-pptx is not installed: pip install python-pptx")
    with tempfile.TemporaryDirectory(prefix="pptx-build-") as work:
        tpl = load_template(a.template, work)
        try:
            prs = Presentation(tpl)
        except Exception as exc:
            raise pk.Unsupported("python-pptx cannot open the template: %s" % pk.clean_text(str(exc), 200))
        lays = layouts_of(prs)
        if a.list_layouts:
            return {"ok": True, "template": a.template,
                    "slide_size_emu": {"cx": prs.slide_width, "cy": prs.slide_height},
                    "template_slides": len(prs.slides),
                    "layouts": [{k: v for k, v in l.items() if k != "_obj"} for l in lays], "error": None}
        try:
            with open(a.spec, encoding="utf-8") as fh:
                spec = json.load(fh)
        except (OSError, ValueError) as exc:
            raise pk.Unsupported("cannot read spec: %s" % exc)
        base = os.path.dirname(os.path.abspath(a.spec))
        if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list) or not spec["slides"]:
            raise pk.Unsupported('spec must be {"slides": [ ... ]} with at least one slide')
        old_ids = [s.slide_id for s in prs.slides]
        report = []
        for n, item in enumerate(spec["slides"], 1):
            lay_key = item.get("layout", 1)
            if isinstance(lay_key, int):
                if not 0 <= lay_key < len(lays):
                    raise SpecError("slide %d: layout index %d not in 0..%d" % (n, lay_key, len(lays) - 1))
                lay = lays[lay_key]
            else:
                lay = next((l for l in lays if l["name"].lower() == str(lay_key).lower()), None)
                if lay is None:
                    raise SpecError("slide %d: no layout named %r; have %s" % (n, lay_key, [l["name"] for l in lays]))
            slide = prs.slides.add_slide(lay["_obj"])
            have = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
            filled = []
            for key, value in (item.get("fill") or {}).items():
                try:
                    idx = int(key)
                except ValueError:
                    raise pk.Unsupported("slide %d: fill keys are placeholder idx numbers, got %r" % (n, key))
                if idx not in have:
                    raise SpecError("slide %d (%s): no placeholder idx %d; have %s" % (n, lay["name"], idx, sorted(have)))
                ph = have[idx]
                if isinstance(value, dict) and "image" in value:
                    img = image_path(base, value["image"], a.allow_image_outside, n)
                    if not hasattr(ph, "insert_picture"):
                        raise SpecError("slide %d: placeholder idx %d is not a picture placeholder" % (n, idx))
                    try:
                        ph.insert_picture(img)
                    except Exception as exc:  # PIL and python-pptx raise many types for bad images
                        raise pk.Unsupported("slide %d: cannot use image %s: %s"
                                             % (n, value["image"], pk.clean_text(str(exc), 200)))
                elif isinstance(value, dict) and "table" in value:
                    fill_table(slide, ph, value["table"], "slide %d idx %d" % (n, idx))
                else:
                    if not ph.has_text_frame:
                        raise SpecError("slide %d: placeholder idx %d holds no text" % (n, idx))
                    fill_text(ph.text_frame, value, "slide %d idx %d" % (n, idx))
                filled.append(idx)
            empty = sorted(set(have) - set(filled))
            if a.remove_empty:
                for idx in empty:
                    el = have[idx]._element
                    el.getparent().remove(el)
            if item.get("notes"):
                check(item["notes"], "slide %d notes" % n)
                slide.notes_slide.notes_text_frame.text = item["notes"]
            if item.get("hidden"):
                slide._element.set("show", "0")
            report.append({"index": n, "layout": lay["name"], "filled": sorted(filled),
                           "empty": [] if a.remove_empty else empty, "hidden": bool(item.get("hidden"))})
        stage = os.path.join(work, "stage.pptx")
        try:
            prs.save(stage)
        except Exception as exc:
            raise pk.Unsupported("python-pptx could not save the deck: %s" % pk.clean_text(str(exc), 200))
        pkg = pk.Package(stage)
        pkg.protect = (a.template, a.spec)  # save() refuses these as OUT even with --force
        log = []
        if not a.keep_template_slides and old_ids:
            gone = [s for s in pkg.slides() if s["id"] in {str(i) for i in old_ids}]
            before = ps.reachable(pkg)
            ps.remove_slides(pkg, gone, log, label="template slide")
            ps.collect_garbage(pkg, before, log)
            # Template sections now hold no slides and the new slides sit in none: drop them.
            root = pkg.xml(pkg.main)
            for ext in root.iter("{%s}ext" % pk.NS["p"]):
                if ext.get("uri") == pk.SECTION_EXT:
                    ext.getparent().remove(ext)
                    pkg.set_xml(pkg.main, root)
                    log.append("removed template sections")
                    break
        pkg.save(a.output, force=a.force)
        return {"ok": True, "template": a.template, "output": a.output, "slides": report,
                "removed_template_slides": len(old_ids) if not a.keep_template_slides else 0,
                "next": "Run inspect_pptx.py OUT --json and render it to check overflow.", "error": None}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("template")
    ap.add_argument("spec", nargs="?")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--list-layouts", action="store_true")
    ap.add_argument("--keep-template-slides", action="store_true")
    ap.add_argument("--remove-empty", action="store_true")
    ap.add_argument("--allow-image-outside", action="store_true",
                    help="allow image paths outside the spec folder (ask the user first)")
    ap.add_argument("--force", action="store_true", help="replace OUT if it exists (no backup)")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    try:
        if not a.list_layouts:
            if not a.spec or not a.output:
                raise pk.Unsupported("give TEMPLATE SPEC.json OUT.pptx, or --list-layouts TEMPLATE")
            if not a.output.lower().endswith(".pptx"):
                raise pk.Unsupported("OUT must end in .pptx")
            pk.check_output(a.output, a.template, a.spec)
            if os.path.exists(a.output) and not a.force:
                raise pk.Unsupported("output exists: %s" % a.output)
        res = build(a)
    except SpecError as exc:
        res = {"ok": False, "error": pk.clean_text(str(exc), 500), "exit": 1}
        pk.emit(res, a.json, lambda r: print("error: " + r["error"]))
        return 1
    except (pk.Unsupported, pk.Refused) as exc:
        return pk.fail(exc, a.json)

    def human(r):
        if "layouts" in r:
            for l in r["layouts"]:
                print("%2d  %-28s %s" % (l["index"], l["name"], ", ".join(
                    "idx %s %s" % (p["idx"], p["type"]) for p in l["placeholders"])))
        else:
            print("wrote %s (%d slides)" % (r["output"], len(r["slides"])))
    pk.emit(res, a.json, human)
    return 0


if __name__ == "__main__":
    sys.exit(main())
