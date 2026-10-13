#!/usr/bin/env python3
"""Rebuild tests/fixtures/*.pptx. Developer tool only: needs python-pptx and XlsxWriter.

Writes to a temp folder unless --write is given. The smoke tests use a copy of the committed fixtures and never need python-pptx. Run this only when a
fixture must change, then inspect the result and commit it.

deck.pptx (4 slides, two sections "Intro" and "Data"):
  1  title "Q3 revenue" split into runs "Q3 rev" + "enue" (bold) + " review"; body "Q3 revenue
     grew"; speaker note "Q3 revenue talking point"; one legacy comment; a link that jumps to hidden slide 2
  2  hidden slide (show="0") with "Q3 revenue draft"; note "draft note"
  3  column chart with an embedded workbook, category "Q3 revenue"
  4  "Q3" + line break + "revenue", and a slide-number field
core.xml creator "Jane Secret", app.xml Company "Contoso Secret Projects".
"""
import os
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

HERE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def build(tmp):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    para = s1.shapes.title.text_frame.paragraphs[0]
    for text, bold in (("Q3 rev", None), ("enue", True), (" review", None)):
        r = para.add_run()
        r.text = text
        if bold:
            r.font.bold = True
    s1.placeholders[1].text_frame.text = "Q3 revenue grew"
    s1.notes_slide.notes_text_frame.text = "Q3 revenue talking point"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Q3 revenue draft"
    s2.notes_slide.notes_text_frame.text = "draft note"
    s2._element.set("show", "0")
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Chart"
    cd = CategoryChartData()
    cd.categories = ["Q3 revenue", "Q4"]
    cd.add_series("Sales", (1.0, 2.0))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Ending"
    tb = s4.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    tb.text_frame.text = "Q3\vrevenue"  # python-pptx writes \v as a:br
    prs.core_properties.author = "Jane Secret"
    prs.core_properties.last_modified_by = "Jane Secret"
    prs.save(tmp)


def patch(tmp, out):
    zin = zipfile.ZipFile(tmp)
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        d = zin.read(info.filename)
        n = info.filename
        if n == "ppt/presentation.xml":
            sec = ('<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
                   '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
                   '<p14:section name="Intro" id="{11111111-1111-1111-1111-111111111111}"><p14:sldIdLst>'
                   '<p14:sldId id="256"/><p14:sldId id="257"/></p14:sldIdLst></p14:section>'
                   '<p14:section name="Data" id="{22222222-2222-2222-2222-222222222222}"><p14:sldIdLst>'
                   '<p14:sldId id="258"/><p14:sldId id="259"/></p14:sldIdLst></p14:section>'
                   '</p14:sectionLst></p:ext></p:extLst>')
            d = d.replace(b"</p:presentation>", sec.encode() + b"</p:presentation>")
        elif n == "ppt/slides/slide4.xml":
            fld = ('<a:fld id="{33333333-3333-3333-3333-333333333333}" type="slidenum"><a:rPr lang="en-US"/>'
                   '<a:t>4</a:t></a:fld>')
            d = d.replace(b"<a:t>revenue</a:t></a:r>", b"<a:t>revenue</a:t></a:r>" + fld.encode(), 1)
        elif n == "ppt/slides/_rels/slide1.xml.rels":
            d = d.replace(b"</Relationships>",
                          b'<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="../comments/comment1.xml"/>'
                          b'<Relationship Id="rIdJ2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slide2.xml"/></Relationships>')
        elif n == "ppt/slides/slide1.xml":
            d = d.replace(b'<a:r><a:t>Q3 revenue grew</a:t>',
                          b'<a:r><a:rPr lang="en-US" dirty="0"><a:hlinkClick r:id="rIdJ2" action="ppaction://hlinksldjump"/></a:rPr><a:t>Q3 revenue grew</a:t>')
        elif n == "[Content_Types].xml":
            d = d.replace(b"</Types>",
                          b'<Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/>'
                          b'<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/></Types>')
        elif n == "ppt/_rels/presentation.xml.rels":
            d = d.replace(b"</Relationships>",
                          b'<Relationship Id="rIdCA" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors" Target="commentAuthors.xml"/></Relationships>')
        elif n == "docProps/app.xml":
            d = d.replace(b"</Properties>", b"<Company>Contoso Secret Projects</Company></Properties>")
        zout.writestr(info, d)
    zout.writestr("ppt/comments/comment1.xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<p:cmLst xmlns:p="%s"><p:cm authorId="0" dt="2026-10-01T10:00:00.000" idx="1">'
                  '<p:pos x="10" y="10"/><p:text>Cut the Q3 number before sending</p:text></p:cm></p:cmLst>' % P)
    zout.writestr("ppt/commentAuthors.xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<p:cmAuthorLst xmlns:p="%s"><p:cmAuthor id="0" name="Jane Secret" initials="JS" lastIdx="1" clrIdx="0"/>'
                  '</p:cmAuthorLst>' % P)
    zout.close()


if __name__ == "__main__":
    # Writes to a temp folder unless --write is given, so a stray run cannot replace the
    # tracked fixture. Compare the temp output with inspect_pptx.py first.
    import sys
    import tempfile
    dest = HERE if "--write" in sys.argv else tempfile.mkdtemp(prefix="pptx-fixture-")
    os.makedirs(dest, exist_ok=True)
    tmp = os.path.join(dest, ".tmp.pptx")
    build(tmp)
    out = os.path.join(dest, "deck.pptx")
    if os.path.exists(out):
        os.unlink(out)
    patch(tmp, out)
    os.unlink(tmp)
    print("wrote", out, "" if dest == HERE else "(temp; pass --write to replace the tracked fixture)")
