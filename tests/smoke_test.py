#!/usr/bin/env python3
"""Smoke tests for the microsoft-powerpoint skill. Run: python tests/smoke_test.py [-v]

Uses the committed fixture tests/fixtures/deck.pptx and builds threat decks in a temp folder
by editing its XML. Needs only Python 3.9+ and lxml. Render tests run when LibreOffice,
bubblewrap and pdftoppm are present and print SKIP otherwise. python-pptx is used, when
installed, only to confirm that outputs still open.

Prints PASS / FAIL / SKIP per test and a summary; exit 1 if anything failed.
"""
import hashlib
import http.server
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "scripts")
TRACKED = os.path.join(ROOT, "tests", "fixtures", "deck.pptx")
sys.path.insert(0, SCRIPTS)
VERBOSE = "-v" in sys.argv
RESULTS = []
TMP = tempfile.mkdtemp(prefix="pptx-skill-test-")
# Tests read a temp copy, so no test can change the tracked fixture; the last test checks it.
FIX = os.path.join(TMP, "deck.pptx")
shutil.copyfile(TRACKED, FIX)
TRACKED_SHA = __import__("hashlib").sha256(open(TRACKED, "rb").read()).hexdigest()


def run(script, *args):
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args),
                       capture_output=True, text=True, timeout=300)
    data = None
    if "--json" in args:
        try:
            data = json.loads(p.stdout)
        except ValueError:
            pass
    return p.returncode, data, p.stdout + p.stderr


def test(name):
    def deco(fn):
        def wrapped():
            try:
                r = fn()
                status = "SKIP" if isinstance(r, str) else "PASS"
                RESULTS.append((status, name, r if isinstance(r, str) else ""))
            except AssertionError as exc:
                RESULTS.append(("FAIL", name, str(exc)))
            except Exception as exc:  # noqa
                RESULTS.append(("FAIL", name, "%s: %s" % (type(exc).__name__, exc)))
            print("%s  %s %s" % (RESULTS[-1][0], name, RESULTS[-1][2] if RESULTS[-1][2] else ""))
        wrapped.__name__ = fn.__name__
        TESTS.append(wrapped)
        return wrapped
    return deco


TESTS = []


def tmp(name):
    return os.path.join(TMP, name)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def mutate(out, edits=None, add=None, drop=(), dup=None):
    """Copy the fixture with byte replacements per part, extra parts and dropped parts."""
    edits, add = edits or {}, add or {}
    zin = zipfile.ZipFile(FIX)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            if info.filename in drop:
                continue
            d = zin.read(info.filename)
            for old, new in edits.get(info.filename, []):
                assert old in d, "fixture edit anchor missing in %s: %r" % (info.filename, old[:60])
                d = d.replace(old, new, 1)
            z.writestr(info, d)
        for n, d in add.items():
            z.writestr(n, d)
        if dup:
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                z.writestr(dup, zin.read(dup))
    return out


REL = b'<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/%s" Target="%s"%s/>'


def rel(rid, typ, target, external=True):
    return REL % (rid.encode(), typ.encode(), target.encode(), b' TargetMode="External"' if external else b"")


def pic(rid, sid, x=0, attr="r:link"):
    return ('<p:pic><p:nvPicPr><p:cNvPr id="%d" name="Linked %d"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
            '<p:blipFill><a:blip %s="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr><a:xfrm>'
            '<a:off x="%d" y="0"/><a:ext cx="2000000" cy="2000000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
            '</a:prstGeom></p:spPr></p:pic>' % (sid, sid, attr, rid, x)).encode()


def threat_deck(out, port=1, local_png="/nonexistent.png"):
    """Fixture plus: VBA (as .pptm), program action, javascript link, UNC link, OLE object,
    external http and file pictures, a hidden shape and an off-slide shape on slide 4."""
    s4 = "ppt/slides/slide4.xml"
    shapes = (
        pic("rIdH", 40) + pic("rIdF", 41, 3000000)
        + b'<p:sp><p:nvSpPr><p:cNvPr id="42" name="Run calc"><a:hlinkClick r:id="rIdP" action="ppaction://program"/></p:cNvPr><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="100" y="100"/><a:ext cx="100" cy="100"/></a:xfrm></p:spPr></p:sp>'
        + b'<p:sp><p:nvSpPr><p:cNvPr id="43" name="Secret box" hidden="1"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="200" y="200"/><a:ext cx="100" cy="100"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:p><a:r><a:t>hidden words</a:t></a:r></a:p></p:txBody></p:sp>'
        + b'<p:sp><p:nvSpPr><p:cNvPr id="44" name="Parked"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="99999999" y="200"/><a:ext cx="100" cy="100"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:p><a:r><a:rPr><a:hlinkClick r:id="rIdJ"/></a:rPr><a:t>click</a:t></a:r><a:r><a:rPr><a:hlinkClick r:id="rIdU"/></a:rPr><a:t>share</a:t></a:r></a:p></p:txBody></p:sp>'
        + b'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="45" name="Object 1"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="0" y="3000000"/><a:ext cx="1000000" cy="1000000"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole"><p:oleObj r:id="rIdO" progId="Package" showAsIcon="1"><p:embed/></p:oleObj></a:graphicData></a:graphic></p:graphicFrame>'
    )
    rels = (rel("rIdH", "image", "http://127.0.0.1:%d/canary.png" % port)
            + rel("rIdF", "image", "file://" + local_png)
            + rel("rIdP", "hyperlink", "file:///C:/Windows/System32/calc.exe")
            + rel("rIdJ", "hyperlink", "javascript:alert(1)")
            + rel("rIdU", "hyperlink", "\\\\evil-host\\share\\x.pptx")
            + rel("rIdO", "oleObject", "../embeddings/oleObject1.bin", external=False))
    return mutate(out, edits={
        s4: [(b"</p:spTree>", shapes + b"</p:spTree>")],
        "ppt/slides/_rels/slide4.xml.rels": [(b"</Relationships>", rels + b"</Relationships>")],
        "[Content_Types].xml": [
            (b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
             b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"),
            (b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')],
        "ppt/_rels/presentation.xml.rels": [(b"</Relationships>",
            b'<Relationship Id="rIdV" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')],
    }, add={"ppt/vbaProject.bin": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 not real vba",
            "ppt/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 payload"})


def opens(path):
    try:
        import pptx  # noqa
    except ImportError:
        return True
    from pptx import Presentation
    Presentation(path)
    return True


def consistent(path):
    """Every internal relationship resolves and every Override names an existing part."""
    import pptpkg as pk
    pkg = pk.Package(path)
    for n in pkg.names:
        if n.endswith(".rels"):
            src = n.replace("_rels/", "")[:-5] if n != "_rels/.rels" else None
            for r in pkg.rels(src):
                if not r["external"]:
                    assert pkg.exists(r["part"]), "dangling rel %s -> %s" % (n, r["part"])
    for o in pkg.ct.findall("ct:Override", pk.NS):
        assert o.get("PartName").lstrip("/") in pkg.names, "Override for missing part %s" % o.get("PartName")
    ids = [s["id"] for s in pkg.slides()]
    for sec in pkg.sections():
        for i in sec["slide_ids"]:
            assert i in ids, "section %s lists missing slide id %s" % (sec["name"], i)
    return True


# ---------------------------------------------------------------- inspect
@test("inspect: slide order, hidden slide, sections, notes, comments, chart workbook, metadata")
def t_inspect():
    rc, d, out = run("inspect_pptx.py", FIX, "--json")
    assert rc == 0 and d["ok"], out
    assert [s["index"] for s in d["slides"]] == [1, 2, 3, 4]
    assert [s["hidden"] for s in d["slides"]] == [False, True, False, False]
    assert [x["name"] for x in d["sections"]] == ["Intro", "Data"]
    assert d["totals"]["notes_with_text"] == 2 and d["totals"]["comments"] == 1
    assert d["slides"][2]["charts"][0]["embedded_workbook"].startswith("ppt/embeddings/")
    assert d["slides"][0]["title"] == "Q3 revenue review"
    assert d["metadata"]["creator"] == "Jane Secret" and d["metadata"]["Company"].startswith("Contoso")
    names = {r["risk"] for r in d["risks"]}
    assert {"hidden-slides", "speaker-notes", "comments", "metadata"} <= names, names


@test("inspect: slide order follows p:sldIdLst, not part names")
def t_order():
    p = mutate(tmp("order.pptx"), edits={"ppt/presentation.xml": [
        (b'<p:sldId id="256" r:id="rId7"/><p:sldId id="257" r:id="rId9"/>',
         b'<p:sldId id="257" r:id="rId9"/><p:sldId id="256" r:id="rId7"/>')]})
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 0, out
    assert d["slides"][0]["part"] == "ppt/slides/slide2.xml", d["slides"][0]


@test("inspect: placeholder inheritance, orphan placeholder reported")
def t_ph():
    p = mutate(tmp("orphan.pptx"), edits={"ppt/slides/slide1.xml": [(b'<p:ph idx="1"/>', b'<p:ph type="pic" idx="77"/>')]})
    rc, d, out = run("inspect_pptx.py", p, "--json")
    phs = d["slides"][0]["placeholders"]
    assert phs[0]["inherits_from"] == "layout", phs
    assert any(x["inherits_from"] == "none" for x in phs), phs
    assert d["totals"]["orphan_placeholders"] == 1


@test("inspect: printed text has bidi, zero-width and control characters removed")
def t_clean():
    p = mutate(tmp("bidi.pptx"), edits={"ppt/slides/slide1.xml": [(b"<a:t>Q3 rev</a:t>", "<a:t>Q3\u202e\u200b rev\x1b</a:t>".replace("\x1b", "").encode())]})
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 0, out
    t = d["slides"][0]["title"]
    assert "\u202e" not in t and "\u200b" not in t, repr(t)


@test("inspect: external targets hidden by default, risky schemes and actions flagged")
def t_links():
    p = threat_deck(tmp("threat-a.pptm"))
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 0, out
    assert "evil-host" not in json.dumps(d["external"]) or all("target" not in e for e in d["external"])
    assert all("target" not in e for e in d["external"])
    rc, d2, _ = run("inspect_pptx.py", p, "--json", "--show-targets")
    assert any("calc.exe" in e.get("target", "") for e in d2["external"])
    acts = d["slides"][3]["actions"]
    risks = " ".join(a.get("risk", "") for a in acts)
    assert "runs a program" in risks and "javascript" in risks and "unc" in risks, acts
    names = {r["risk"] for r in d["risks"]}
    assert {"vba-macros", "ole-objects", "external-link", "hidden-shapes", "offslide-shapes", "action"} <= names, names
    assert d["kind"] == "pptm"


@test("refuse: duplicate zip entry (exit 3)")
def t_dup():
    p = mutate(tmp("dup.pptx"), dup="ppt/slides/slide1.xml")
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 3 and "duplicate" in d["error"], out


@test("refuse: entry names that differ only by case (exit 3)")
def t_case():
    p = mutate(tmp("case.pptx"), add={"PPT/slides/slide1.xml": b"<x/>"})
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 3 and "case" in d["error"], out


@test("refuse: DOCTYPE / entity expansion (billion laughs) without expanding (exit 3)")
def t_doctype():
    bomb = b'<?xml version="1.0"?><!DOCTYPE p [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">&b;</p:sld>'
    zin = zipfile.ZipFile(FIX)
    p = tmp("bomb.pptx")
    with zipfile.ZipFile(p, "w") as z:
        for i in zin.infolist():
            z.writestr(i, bomb if i.filename == "ppt/slides/slide1.xml" else zin.read(i.filename))
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 3 and "DOCTYPE" in d["error"], out


@test("refuse: too many entries and unsafe entry names (exit 3)")
def t_caps():
    import pptpkg as pk
    old = pk.MAX_MEMBERS
    pk.MAX_MEMBERS = 10
    try:
        try:
            pk.Package(FIX)
            raise AssertionError("cap not enforced")
        except pk.Refused as exc:
            assert "too many" in str(exc)
    finally:
        pk.MAX_MEMBERS = old
    p = mutate(tmp("slip.pptx"), add={"../evil.xml": b"<x/>"})
    rc, d, out = run("inspect_pptx.py", p, "--json")
    assert rc == 3 and "unsafe" in d["error"], out


@test("unsupported: encrypted, legacy .ppt and non-zip input give exit 2 with a reason")
def t_unsupported():
    head = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 500
    enc = tmp("enc.pptx")
    open(enc, "wb").write(head + "EncryptionInfo".encode("utf-16-le"))
    rc, d, _ = run("inspect_pptx.py", enc, "--json")
    assert rc == 2 and "password" in d["error"], d
    ppt = tmp("old.ppt")
    open(ppt, "wb").write(head + "PowerPoint Document".encode("utf-16-le"))
    rc, d, _ = run("inspect_pptx.py", ppt, "--json")
    assert rc == 2 and ".ppt" in d["error"], d
    txt = tmp("x.pptx")
    open(txt, "w").write("hello")
    rc, d, _ = run("inspect_pptx.py", txt, "--json")
    assert rc == 2 and d["ok"] is False


# ---------------------------------------------------------------- replace
@test("replace: text split across runs is replaced, run formatting kept, other parts byte-identical")
def t_replace():
    out = tmp("rep.pptx")
    before = sha(FIX)
    rc, d, o = run("pptx_replace.py", FIX, out, "--find", "Q3 revenue", "--replace", "Q4 sales",
                   "--slides", "1", "--expect", "2", "--json")
    assert rc == 0 and d["replaced"] == 2, o
    assert sha(FIX) == before, "input changed"
    z, z0 = zipfile.ZipFile(out), zipfile.ZipFile(FIX)
    x = z.read("ppt/slides/slide1.xml").decode()
    assert "Q4 sales" in x and "Q3 rev" not in x
    assert '<a:rPr b="1"/>' in x or 'b="1"' in x, "bold run lost"
    for n in z0.namelist():
        if n != "ppt/slides/slide1.xml":
            assert z.read(n) == z0.read(n), "part changed: " + n
    assert opens(out)


@test("replace: --expect mismatch writes nothing (exit 1)")
def t_expect():
    out = tmp("rep-expect.pptx")
    rc, d, o = run("pptx_replace.py", FIX, out, "--find", "Q3 revenue", "--replace", "X", "--expect", "1", "--json")
    assert rc == 1 and not os.path.exists(out), o
    assert d["replaced"] == 3, d  # slides 1 (x2) and hidden slide 2


@test("replace: hidden slide reported, notes only with --notes, chart text not replaced, breaks and fields skipped")
def t_scope():
    rc, d, o = run("pptx_replace.py", FIX, tmp("dry.pptx"), "--find", "Q3 revenue", "--replace", "X", "--dry-run", "--json")
    assert rc == 0 and not os.path.exists(tmp("dry.pptx")), o
    assert d["hidden_slides_touched"] == ["slide 2"], d
    assert d["not_replaced"].get("charts", 0) >= 1, d["not_replaced"]
    assert d["skipped_across_breaks_or_fields"] == 1, d  # "Q3" a:br "revenue" on slide 4
    rc, d2, _ = run("pptx_replace.py", FIX, tmp("dry2.pptx"), "--find", "Q3 revenue", "--replace", "X", "--dry-run", "--notes", "--json")
    assert d2["replaced"] == d["replaced"] + 1, (d2["replaced"], d["replaced"])


@test("replace: regex with groups; catastrophic regex killed by timeout (exit 2)")
def t_regex():
    out = tmp("rx.pptx")
    rc, d, o = run("pptx_replace.py", FIX, out, "--find", r"Q(\d) revenue grew", "--replace", r"Quarter \1 up",
                   "--regex", "--json")
    assert rc == 0 and "Quarter 3 up" in zipfile.ZipFile(out).read("ppt/slides/slide1.xml").decode(), o
    p = mutate(tmp("redos.pptx"), edits={"ppt/slides/slide4.xml": [(b"<a:t>Ending</a:t>", b"<a:t>" + b"a" * 40 + b"!</a:t>")]})
    t0 = time.time()
    rc, d, o = run("pptx_replace.py", p, tmp("redos-out.pptx"), "--find", "(a+)+$", "--replace", "x", "--regex",
                   "--timeout", "1", "--json")
    assert rc == 2 and "did not finish" in d["error"], o
    assert time.time() - t0 < 20


@test("replace: refuses an existing output and the input as output")
def t_noover():
    out = tmp("exists.pptx")
    open(out, "w").write("keep")
    rc, d, o = run("pptx_replace.py", FIX, out, "--find", "Ending", "--replace", "End", "--json")
    assert rc == 2 and open(out).read() == "keep", o
    cp = tmp("self.pptx")
    shutil.copy(FIX, cp)
    rc, d, o = run("pptx_replace.py", cp, cp, "--find", "Ending", "--replace", "End", "--json")
    assert rc == 2 and sha(cp) == sha(FIX), o


# ---------------------------------------------------------------- sanitize
@test("share-check: fixture is not clean (exit 1) and lists residue")
def t_check():
    rc, d, o = run("pptx_sanitize.py", FIX, "--check-clean", "--json")
    assert rc == 1 and not d["ok"], o
    assert {"speaker-notes", "comments", "hidden-slides", "metadata"} <= {r["risk"] for r in d["residue"]}


@test("sanitize --all: threat deck becomes clean .pptx, package consistent, opens")
def t_all():
    src = threat_deck(tmp("threat-b.pptm"))
    out = tmp("clean.pptx")
    rc, d, o = run("pptx_sanitize.py", src, out, "--all", "--json")
    assert rc == 0 and d["ok"] and d["kind"] == "pptx", o
    names = zipfile.ZipFile(out).namelist()
    assert not any("vbaProject" in n or "oleObject" in n or n.startswith("ppt/notesSlides/")
                   or n.startswith("ppt/comments/") for n in names), names
    blob = b"".join(zipfile.ZipFile(out).read(n) for n in names if n.endswith((".xml", ".rels")))
    for bad in (b"calc.exe", b"javascript:", b"evil-host", b"canary.png", b"hidden words", b"Jane Secret",
                b"Contoso", b"Q3 revenue draft", b"ppaction://program"):
        assert bad not in blob, bad
    rc, d2, o2 = run("pptx_sanitize.py", out, "--check-clean", "--json")
    assert rc == 0, o2
    assert consistent(out) and opens(out)


@test("sanitize: macro-free output must use .pptx; .pptm output refused (exit 2)")
def t_ext():
    src = threat_deck(tmp("threat-c.pptm"))
    rc, d, o = run("pptx_sanitize.py", src, tmp("still.pptm"), "--remove-macros", "--json")
    assert rc == 2 and ".pptx" in d["error"] and not os.path.exists(tmp("still.pptm")), o
    rc, d, o = run("pptx_sanitize.py", src, tmp("nomacro.pptx"), "--remove-macros", "--json")
    assert rc == 0, o
    assert "ppt/vbaProject.bin" not in zipfile.ZipFile(tmp("nomacro.pptx")).namelist()


@test("sanitize: without --remove-macros a .pptm stays .pptm with its vbaProject.bin unchanged")
def t_keep_macro():
    src = threat_deck(tmp("threat-g.pptm"))
    out = tmp("kept.pptm")
    rc, d, o = run("pptx_sanitize.py", src, out, "--remove-notes", "--remove-comments", "--json")
    assert rc == 0 and d["kind"] == "pptm", o
    assert zipfile.ZipFile(out).read("ppt/vbaProject.bin") == zipfile.ZipFile(src).read("ppt/vbaProject.bin")
    rc, d, o = run("pptx_sanitize.py", out, "--check-clean", "--json")
    assert rc == 1 and "vba-macros" in {r["risk"] for r in d["residue"]}, o


@test("sanitize: removing hidden slides keeps sections and links consistent")
def t_hidden():
    out = tmp("nohidden.pptx")
    rc, d, o = run("pptx_sanitize.py", FIX, out, "--remove-hidden-slides", "--json")
    assert rc == 0, o
    rc, d2, _ = run("inspect_pptx.py", out, "--json")
    assert d2["totals"]["slides"] == 3 and d2["totals"]["hidden"] == 0
    assert [len(s["slide_ids"]) for s in d2["sections"]] == [1, 2]
    assert "hlinkClick" not in zipfile.ZipFile(out).read("ppt/slides/slide1.xml").decode()
    assert consistent(out) and opens(out)


@test("sanitize: slide part left out of the slide list (python-pptx delete leftover) is found and removed")
def t_unlisted():
    p = mutate(tmp("unlisted.pptx"), edits={"ppt/presentation.xml": [(b'<p:sldId id="257" r:id="rId9"/>', b"")]})
    rc, d, o = run("inspect_pptx.py", p, "--json")
    assert d["unlisted_slide_parts"] == ["ppt/slides/slide2.xml"] and d["stale_section_ids"] == ["257"], o
    assert "unlisted-slides" in {r["risk"] for r in d["risks"]}
    out = tmp("unlisted-clean.pptx")
    rc, d, o = run("pptx_sanitize.py", p, out, "--remove-hidden-slides", "--json")
    assert rc == 0, o
    assert "ppt/slides/slide2.xml" not in zipfile.ZipFile(out).namelist()
    assert b"Q3 revenue draft" not in b"".join(zipfile.ZipFile(out).read(n) for n in zipfile.ZipFile(out).namelist())
    assert consistent(out) and opens(out)


@test("sanitize: one category only; residue judged on requested categories; --dry-run writes nothing")
def t_cat():
    out = tmp("nonotes.pptx")
    rc, d, o = run("pptx_sanitize.py", FIX, out, "--remove-notes", "--json")
    assert rc == 0 and d["ok"], o
    rc, d2, _ = run("inspect_pptx.py", out, "--json")
    assert d2["totals"]["notes_with_text"] == 0 and d2["totals"]["comments"] == 1
    rc, d3, o3 = run("pptx_sanitize.py", FIX, tmp("dry-s.pptx"), "--all", "--dry-run", "--json")
    assert rc == 0 and d3["would_remove"] and not os.path.exists(tmp("dry-s.pptx")), o3


@test("sanitize: residue is detected (exit 1) when a requested category is not cleared")
def t_residue():
    import pptx_sanitize as ps
    orig = ps.do_comments
    ps.do_comments = lambda pkg, log: None  # simulate a removal that misses something
    try:
        buf = io.StringIO()
        old = sys.stdout
        sys.stdout = buf
        try:
            rc = ps.main([FIX, tmp("residue.pptx"), "--remove-comments", "--json"])
        finally:
            sys.stdout = old
        d = json.loads(buf.getvalue())
        assert rc == 1 and not d["ok"] and d["remaining"], d
    finally:
        ps.do_comments = orig



# ---------------------------------------------------------------- 0.1.1 security review fixes
def hostile_deck(out):
    """VBA under a neutral name (found by content type and rel), a program action on the master,
    a javascript link and a mismatched link in notes, an external picture in the theme rels."""
    return mutate(out, edits={
        "[Content_Types].xml": [
            (b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
             b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"),
            (b"</Types>", b'<Override PartName="/ppt/stuff/data.dat" ContentType="application/vnd.ms-office.vbaProject"/></Types>')],
        "ppt/_rels/presentation.xml.rels": [(b"</Relationships>",
            b'<Relationship Id="rIdV" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="stuff/data.dat"/></Relationships>')],
        "ppt/slideMasters/slideMaster1.xml": [(b"</p:spTree>",
            b'<p:sp><p:nvSpPr><p:cNvPr id="99" name="M"><a:hlinkClick r:id="rIdMP" action="ppaction://program"/></p:cNvPr><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/></p:sp></p:spTree>')],
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": [(b"</Relationships>",
            rel("rIdMP", "hyperlink", "file:///C:/x.exe") + b"</Relationships>")],
        "ppt/notesSlides/notesSlide1.xml": [(b"<a:r><a:t>Q3 revenue talking point",
            b'<a:r><a:rPr><a:hlinkClick r:id="rIdNJ"/></a:rPr><a:t>www.bank.com</a:t></a:r>'
            b'<a:r><a:rPr><a:hlinkClick r:id="rIdNK"/></a:rPr><a:t>open</a:t></a:r><a:r><a:t>Q3 revenue talking point')],
        "ppt/notesSlides/_rels/notesSlide1.xml.rels": [(b"</Relationships>",
            rel("rIdNJ", "hyperlink", "https://evil.example/login") + rel("rIdNK", "hyperlink", "javascript:alert(1)")
            + b"</Relationships>")],
    }, add={"ppt/stuff/data.dat": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 vba",
            "ppt/theme/_rels/theme1.xml.rels": b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + rel("rIdT", "image", "http://beacon.example/t.png") + b"</Relationships>"})


def all_bytes(path):
    z = zipfile.ZipFile(path)
    return b"".join(z.read(n) for n in z.namelist())


@test("H1: VBA found by content type and relationship, not file name; macro-enabled type flagged")
def t_h1():
    p = hostile_deck(tmp("hostile.pptx"))
    rc, d, o = run("inspect_pptx.py", p, "--json")
    risks = {r["risk"]: r for r in d["risks"]}
    assert risks["vba-macros"]["detail"] == ["ppt/stuff/data.dat"], o
    assert risks["macro-enabled-type"]["detail"] == "pptm"
    out = tmp("hostile-clean.pptx")
    rc, d, o = run("pptx_sanitize.py", p, out, "--remove-macros", "--json")
    assert rc == 0, o
    z = zipfile.ZipFile(out)
    assert "ppt/stuff/data.dat" not in z.namelist() and b"macroEnabled" not in z.read("[Content_Types].xml")
    assert b"vbaProject" not in all_bytes(out) and consistent(out) and opens(out)


@test("H2: actions and risky schemes found in masters and notes, not only slides; removed")
def t_h2():
    p = hostile_deck(tmp("hostile2.pptx"))
    rc, d, o = run("inspect_pptx.py", p, "--json")
    acts = [(r["part"], r["detail"]) for r in d["risks"] if r["risk"] == "action"]
    assert ("ppt/slideMasters/slideMaster1.xml", "runs a program when clicked") in acts, acts
    assert ("ppt/notesSlides/notesSlide1.xml", "link scheme javascript") in acts, acts
    assert "notes of slide 1" in [r["where"] for r in d["risks"] if r["risk"] == "action"]
    out = tmp("hostile2-clean.pptm")
    rc, d, o = run("pptx_sanitize.py", p, out, "--remove-active", "--json")
    assert rc == 0, o
    b = all_bytes(out)
    assert b"javascript:" not in b and b"ppaction://program" not in b and b"x.exe" not in b
    assert b"Q3 revenue talking point" in b, "link removal must keep the notes text"
    assert consistent(out) and opens(out)


@test("M1: external targets in any .rels (theme) found and removed")
def t_m1():
    p = hostile_deck(tmp("hostile3.pptx"))
    rc, d, o = run("inspect_pptx.py", p, "--json")
    assert any(e["host"] == "beacon.example" and e["part"] == "ppt/theme/theme1.xml" for e in d["external"]), o
    out = tmp("hostile3-clean.pptm")
    rc, d, o = run("pptx_sanitize.py", p, out, "--remove-active", "--json")
    assert rc == 0 and b"beacon.example" not in all_bytes(out), o
    assert consistent(out)


@test("M2: link text vs target: other domain, userinfo, no web host, invisible chars; same site passes")
def t_m2():
    import pptlinks as lk
    cases = [("www.bank.com", "https://evil.example/login", "different_domain"),
             ("https://bank.com", "https://bank.com@evil.example/", "userinfo"),
             ("bank.com", "file:///C:/x.exe", "no_web_host"),
             ("bank\u200b.com", "https://bank.com/", "invisible_chars"),
             ("www.bank.com", "https://login.bank.com/x", None),
             ("finance.gov.pk", "https://evil.gov.pk/", "different_domain"),
             ("Click here", "https://anything.example/", None)]
    for text, target, want in cases:
        got = lk.mismatch(text, target)
        assert got == want, (text, target, got, want)
    assert lk.classify("javascript:alert(1)")[2] and lk.classify("data:text/html,x")[2]
    assert lk.classify("\\\\host\\share")[2] and lk.classify("../x.pptx")[2]
    assert not lk.classify("https://example.com")[2] and not lk.classify("mailto:a@b.c")[2]
    p = hostile_deck(tmp("hostile4.pptx"))
    rc, d, o = run("inspect_pptx.py", p, "--json")
    mm = [r for r in d["risks"] if r["risk"] == "mismatched-link"]
    assert mm and mm[0]["text"] == "www.bank.com" and mm[0]["detail"] == "different_domain", o
    rc, d, o = run("pptx_sanitize.py", p, "--check-clean", "--json")
    assert rc == 1 and "mismatched-link" in str(d["residue"]), o


@test("L1: control, bidi and zero-width characters refused in replacement text (exit 2)")
def t_l1():
    for bad in ("evil\u202egnp", "a\x01b", "zero\u200bwidth"):
        out = tmp("l1.pptx")
        rc, d, o = run("pptx_replace.py", FIX, out, "--find", "Q3", "--replace", bad, "--json")
        assert rc == 2 and not os.path.exists(out), (bad, o)
    rc, d, o = run("pptx_replace.py", FIX, tmp("l1b.pptx"), "--find", "(Q3)", "--replace", "\\1\u202e",
                   "--regex", "--json")
    assert rc == 2, o


@test("L2: a bad pattern or child error is reported as such, not as a timeout")
def t_l2():
    t0 = time.time()
    rc, d, o = run("pptx_replace.py", FIX, tmp("l2.pptx"), "--find", "(Q3)", "--replace", "\\9", "--regex", "--json")
    assert rc == 2 and "timeout" not in d["error"].lower() and "timed out" not in d["error"].lower(), o
    assert time.time() - t0 < 4, "error should be immediate"


@test("L3: metadata --check-clean and --strip-metadata cover the same fields incl title, revision, Template, TotalTime")
def t_l3():
    import pptpkg as pk
    for f in ("title", "revision", "creator", "lastModifiedBy"):
        assert f in pk.CORE_STRIP, f
    for f in ("Template", "TotalTime", "Company", "Manager", "HyperlinkBase"):
        assert f in pk.APP_STRIP, f
    p = mutate(tmp("meta.pptx"), edits={
        "docProps/core.xml": [(b"<dc:title/>", b"<dc:title>Secret plan</dc:title>")],
        "docProps/app.xml": [(b"</Properties>", b"<Template>C:\\Users\\jane\\t.potx</Template></Properties>")]})
    rc, d, o = run("inspect_pptx.py", p, "--json")
    meta = [r for r in d["risks"] if r["risk"] == "metadata"][0]["detail"]
    assert {"title", "revision", "TotalTime", "Template"} <= set(meta), meta
    out = tmp("meta-clean.pptx")
    rc, d, o = run("pptx_sanitize.py", p, out, "--strip-metadata", "--json")
    assert rc == 0, o
    rc, d, o = run("inspect_pptx.py", out, "--json")
    assert not [r for r in d["risks"] if r["risk"] == "metadata"], o
    z = zipfile.ZipFile(out)
    b = b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))  # comment authors are --remove-comments
    for secret in (b"Secret plan", b"jane", b"Jane Secret", b"Contoso"):
        assert secret not in b, secret
    rc, d, o = run("pptx_sanitize.py", out, "--check-clean", "--json")
    assert "metadata" not in str(d.get("residue")), o


@test("L4: --check-clean and write results state clean_scope and that ok is not a safety verdict")
def t_l4():
    rc, d, o = run("pptx_sanitize.py", FIX, "--check-clean", "--json")
    sc = d["clean_scope"]
    assert sc["checked"] and any("visible slide text" in x for x in sc["not_checked"]), o
    assert "not that the deck is safe" in d["note"]
    rc, d, o = run("pptx_sanitize.py", FIX, tmp("l4.pptx"), "--all", "--json")
    assert rc == 0 and d["clean_scope"]["not_checked"], o


@test("L6: render checks content not extension; zip renamed .ppt and odp without content.xml refused; memory cap")
def t_l6():
    import pptx_render as pr
    ppt = tmp("renamed.ppt")
    shutil.copyfile(FIX, ppt)
    rc, d, o = run("pptx_render.py", ppt, tmp("renamed.pdf"), "--json")
    assert rc == 2 and not os.path.exists(tmp("renamed.pdf")), o
    odp = tmp("fake.odp")
    with zipfile.ZipFile(odp, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.presentation")
    rc, d, o = run("pptx_render.py", odp, tmp("fake.pdf"), "--json")
    assert rc in (2, 3) and not os.path.exists(tmp("fake.pdf")), o
    bomb = tmp("bomb.odp")
    with zipfile.ZipFile(bomb, "w") as z:
        z.writestr("content.xml", "<a/>")
        for i in range(12000):
            z.writestr("x/%d" % i, "")
    rc, d, o = run("pptx_render.py", bomb, tmp("bomb.pdf"), "--json")
    assert rc == 3, o
    rc, d, o = run("pptx_render.py", FIX, tmp("m.pdf"), "--memory-mb", "100", "--json")
    assert rc == 2 and "memory" in d["error"], o
    assert pr.MEMORY_MB >= 512


# ---------------------------------------------------------------- pptx_slides.py
def order(path):
    rc, d, o = run("inspect_pptx.py", path, "--json")
    return [s["title"] for s in d["slides"]], d


@test("slides: delete, move, hide, unhide; sections and custom shows stay consistent; opens")
def t_slides():
    out = tmp("sl.pptx")
    rc, d, o = run("pptx_slides.py", FIX, out, "--delete", "2", "--move", "3:1", "--expect-slides", "3", "--json")
    assert rc == 0, o
    ORIG = order(FIX)[0]
    titles, d = order(out)
    assert len(titles) == 3 and titles[0] == ORIG[3], (titles, ORIG)
    assert b"Q3 revenue draft" not in all_bytes(out), "deleted slide text must be gone from the package"
    assert consistent(out) and opens(out)
    assert sum(len(s["slide_ids"]) for s in d["sections"]) == 3
    out2 = tmp("sl2.pptx")
    rc, _, o = run("pptx_slides.py", FIX, out2, "--hide", "1", "--unhide", "2", "--json")
    assert rc == 0, o
    _, d2 = order(out2)
    assert [s["hidden"] for s in d2["slides"]][:2] == [True, False], d2["slides"]
    cs = mutate(tmp("custshow.pptx"), edits={"ppt/presentation.xml": [(b"<p:defaultTextStyle>",
        b'<p:custShowLst><p:custShow name="Short" id="0"><p:sldLst><p:sld r:id="rId9"/><p:sld r:id="rId11"/>'
        b'</p:sldLst></p:custShow></p:custShowLst><p:defaultTextStyle>')]})
    out3 = tmp("custshow-out.pptx")
    rc, _, o = run("pptx_slides.py", cs, out3, "--delete", "2", "--json")
    assert rc == 0, o
    pres = zipfile.ZipFile(out3).read("ppt/presentation.xml")
    assert b'r:id="rId9"' not in pres and b'r:id="rId11"' in pres, "custom show entry for the deleted slide must go"
    assert consistent(out3) and opens(out3)


@test("slides: refuses deleting every slide, bad numbers, --expect-slides mismatch, wrong extension; --dry-run writes nothing")
def t_slides_refuse():
    rc, d, o = run("pptx_slides.py", FIX, tmp("none.pptx"), "--delete", "1-4", "--json")
    assert rc == 2 and not os.path.exists(tmp("none.pptx")), o
    rc, d, o = run("pptx_slides.py", FIX, tmp("bad.pptx"), "--delete", "9", "--json")
    assert rc == 2 and not os.path.exists(tmp("bad.pptx")), o
    rc, d, o = run("pptx_slides.py", FIX, tmp("exp.pptx"), "--delete", "1", "--expect-slides", "7", "--json")
    assert rc == 1 and not os.path.exists(tmp("exp.pptx")), o
    for script, args in (("pptx_slides.py", ["--hide", "1"]), ("pptx_replace.py", ["--find", "Q3", "--replace", "Q4"])):
        rc, d, o = run(script, FIX, tmp("wrong.pptm"), *args, "--json")
        assert rc == 2 and "must end in .pptx" in d["error"] and not os.path.exists(tmp("wrong.pptm")), o
    rc, d, o = run("pptx_slides.py", FIX, tmp("dry.pptx"), "--delete", "1", "--dry-run", "--json")
    assert rc == 0 and not os.path.exists(tmp("dry.pptx")), o


# ---------------------------------------------------------------- pptx_build.py
def have_pptx():
    try:
        import pptx  # noqa
        return True
    except ImportError:
        return False


def potx(out):
    zin = zipfile.ZipFile(FIX)
    with zipfile.ZipFile(out, "w") as z:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "[Content_Types].xml":
                d = d.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            z.writestr(i, d)
    return out


@test("build: --list-layouts gives placeholder idx; spec fills by idx from a .potx; template slides and sections removed")
def t_build():
    if not have_pptx():
        return "python-pptx not installed"
    tpl = potx(tmp("t.potx"))
    rc, d, o = run("pptx_build.py", "--list-layouts", tpl, "--json")
    assert rc == 0 and d["layouts"][0]["name"] == "Title Slide", o
    assert 0 in [p["idx"] for p in d["layouts"][1]["placeholders"]]
    red_png(tmp("pic.png"))
    spec = tmp("spec.json")
    json.dump({"slides": [
        {"layout": "Title Slide", "fill": {"0": "Quarterly review", "1": "October 2026"}},
        {"layout": 1, "fill": {"0": "Results", "1": ["Revenue up", {"text": "EMEA led", "level": 1}]}, "notes": "Say July."},
        {"layout": "Picture with Caption", "fill": {"0": "Site", "1": {"image": "pic.png"}}, "hidden": True}]},
        open(spec, "w"))
    out = tmp("built.pptx")
    rc, d, o = run("pptx_build.py", tpl, spec, out, "--json")
    assert rc == 0 and d["slides"][2]["empty"] == [2], o
    titles, d2 = order(out)
    assert titles == ["Quarterly review", "Results", "Site"], titles
    assert d2["slides"][2]["hidden"] and d2["sections"] == []
    assert b"Q3 revenue" not in all_bytes(out), "template slide text must be gone"
    from pptx import Presentation
    s = Presentation(out).slides[1]
    body = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
    assert [p.level for p in body.text_frame.paragraphs] == [0, 1]
    assert s.notes_slide.notes_text_frame.text == "Say July."
    assert consistent(out)


@test("build: unknown placeholder idx or layout exits 1 and writes nothing; unsafe text exits 2")
def t_build_refuse():
    if not have_pptx():
        return "python-pptx not installed"
    for i, sl in enumerate(({"layout": 1, "fill": {"7": "x"}}, {"layout": "No such layout"}, {"layout": 99})):
        spec = tmp("bad%d.json" % i)
        json.dump({"slides": [sl]}, open(spec, "w"))
        rc, d, o = run("pptx_build.py", FIX, spec, tmp("bad%d.pptx" % i), "--json")
        assert rc == 1 and not os.path.exists(tmp("bad%d.pptx" % i)), o
    spec = tmp("bidi.json")
    json.dump({"slides": [{"layout": 1, "fill": {"0": "a\u202eb"}}]}, open(spec, "w"))
    rc, d, o = run("pptx_build.py", FIX, spec, tmp("build-bidi.pptx"), "--json")
    assert rc == 2 and not os.path.exists(tmp("build-bidi.pptx")), o


# ---------------------------------------------------------------- render
def render_ready():
    import pptx_render as pr
    c = pr.check()
    return c if c["ok"] and c["can_png"] else None


class Canary(http.server.BaseHTTPRequestHandler):
    hits = []

    def do_GET(self):
        Canary.hits.append(self.path)
        self.send_response(404)
        self.end_headers()

    do_OPTIONS = do_HEAD = do_GET

    def log_message(self, *a):
        pass


def canary_server():
    srv = http.server.HTTPServer(("127.0.0.1", 0), Canary)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def red_png(path):
    import struct
    import zlib
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 8 for _ in range(8))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    open(path, "wb").write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
                           + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def pdf_images(pdf):
    if not shutil.which("pdfimages"):
        return None
    out = subprocess.run(["pdfimages", "-list", pdf], capture_output=True, text=True).stdout
    return len(out.strip().splitlines()) - 2


@test("threat is real: unsandboxed LibreOffice fetches http and embeds local linked pictures")
def t_threat_real():
    if not render_ready():
        return "LibreOffice, bwrap or pdftoppm missing"
    import pptx_render as pr
    srv = canary_server()
    Canary.hits.clear()
    secret = tmp("secret.png")
    red_png(secret)
    src = threat_deck(tmp("threat-d.pptx"), port=srv.server_address[1], local_png=secret)
    # threat_deck writes a macro content type; make it a plain .pptx for LibreOffice
    fixed = mutate(tmp("threat-e.pptx"))
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(fixed, "w") as z:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "[Content_Types].xml":
                d = d.replace(b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
                              b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")
            z.writestr(i, d)
    prof = tempfile.mkdtemp(prefix="lo-unsafe-")
    outdir = tempfile.mkdtemp(prefix="lo-out-")
    subprocess.run([pr.find_soffice(), "-env:UserInstallation=file://" + prof, "--headless",
                    "--convert-to", "pdf", "--outdir", outdir, fixed], capture_output=True, timeout=180)
    srv.shutdown()
    pdf = os.path.join(outdir, "threat-e.pdf")
    assert os.path.exists(pdf), "LibreOffice produced no PDF"
    assert Canary.hits, "expected the unsandboxed run to contact the canary server"
    n = pdf_images(pdf)
    assert n is None or n >= 1, "expected the local picture to be embedded"
    t_threat_real.fixed = fixed


@test("render: sandboxed run makes no network request and embeds no local file")
def t_render_sandbox():
    if not render_ready():
        return "LibreOffice, bwrap or pdftoppm missing"
    srv = canary_server()
    Canary.hits.clear()
    secret = tmp("secret2.png")
    red_png(secret)
    src = threat_deck(tmp("threat-f.pptm"), port=srv.server_address[1], local_png=secret)
    out = tmp("boxed.pdf")
    rc, d, o = run("pptx_render.py", src, out, "--json")
    srv.shutdown()
    assert rc == 0 and os.path.exists(out), o
    assert not Canary.hits, Canary.hits
    n = pdf_images(out)
    assert n in (None, 0), "local picture leaked into the PDF (%s images)" % n


@test("render: PDF leaves out hidden slides; --notes-pages; PNG one file per visible slide")
def t_render_out():
    if not render_ready():
        return "LibreOffice, bwrap or pdftoppm missing"
    out = tmp("deck.pdf")
    rc, d, o = run("pptx_render.py", FIX, out, "--json")
    assert rc == 0, o
    txt = subprocess.run(["pdftotext", out, "-"], capture_output=True, text=True).stdout
    assert "Q3 revenue draft" not in txt and "Ending" in txt, txt
    rc, d, o = run("pptx_render.py", FIX, tmp("notes.pdf"), "--notes-pages", "--json")
    txt = subprocess.run(["pdftotext", tmp("notes.pdf"), "-"], capture_output=True, text=True).stdout
    assert rc == 0 and "talking point" in txt, o
    rc, d, o = run("pptx_render.py", FIX, tmp("pngs"), "--to", "png", "--json")
    assert rc == 0 and len(d["files"]) == 3, o


@test("render: refusals (existing output, wrong extension, no bwrap) and timeout (exit 3)")
def t_render_refuse():
    import pptx_render as pr
    out = tmp("exists.pdf")
    open(out, "w").write("keep")
    rc, d, o = run("pptx_render.py", FIX, out, "--json")
    assert rc == 2 and open(out).read() == "keep", o
    rc, d, o = run("pptx_render.py", FIX, tmp("x.png"), "--json")
    assert rc == 2 and ".pdf" in d["error"], o
    real = pr.shutil.which
    pr.shutil.which = lambda n: None if n == "bwrap" else real(n)
    buf, old = io.StringIO(), sys.stdout
    sys.stdout = buf
    try:
        rc = pr.main([FIX, tmp("nobwrap.pdf"), "--json"])
    finally:
        sys.stdout = old
        pr.shutil.which = real
    assert rc == 2 and "sandbox" in json.loads(buf.getvalue())["error"] and not os.path.exists(tmp("nobwrap.pdf"))
    if not render_ready():
        return "timeout part skipped: LibreOffice or bwrap missing"
    rc, d, o = run("pptx_render.py", FIX, tmp("slow.pdf"), "--timeout", "0.2", "--json")
    assert rc == 3 and not os.path.exists(tmp("slow.pdf")), o


# ---------------------------------------------------------------- release
@test("release: skill.json, package.json, SKILL.md and CHANGELOG agree; files exist")
def t_release():
    pj = json.load(open(os.path.join(ROOT, "package.json")))
    sj = json.load(open(os.path.join(ROOT, "skill.json")))
    skill = open(os.path.join(ROOT, "SKILL.md")).read()
    fm = skill.split("---")[1]
    assert pj["name"] == sj["name"] == "microsoft-powerpoint"
    assert re.search(r'^name: microsoft-powerpoint$', fm, re.M)
    v = pj["version"]
    assert sj["version"] == v and ('version: "%s"' % v) in fm
    assert "## %s" % v in open(os.path.join(ROOT, "CHANGELOG.md")).read() or "## [%s]" % v in open(os.path.join(ROOT, "CHANGELOG.md")).read()
    assert sj["description"] == re.search(r'^description: "(.*)"$', fm, re.M).group(1)
    assert sj["description"].startswith("Use when") and len(sj["description"]) <= 1024
    assert len(skill.splitlines()) < 500
    assert "scripts" not in pj or not any(k in pj.get("scripts", {}) for k in ("install", "postinstall", "preinstall"))
    for f in pj["files"]:
        assert os.path.exists(os.path.join(ROOT, f.split("*")[0])), f
    for s in ("inspect_pptx.py", "pptx_sanitize.py", "pptx_replace.py", "pptx_render.py", "pptx_slides.py",
              "pptx_build.py"):
        assert s in skill, s + " not in SKILL.md"
        assert "Exit codes" in open(os.path.join(SCRIPTS, s)).read()


@test("docs: automation.md keeps the unsaved-path guard, PathSeparator, named PDF args and curl -L note")
def t_docs_automation():
    t = open(os.path.join(ROOT, "references", "automation.md"), encoding="utf-8").read()
    for needle in ('ActivePresentation.Path = ""', "Application.PathSeparator", "PrintHiddenSlides:=msoFalse",
                   "curl -sS -L", "ppPlaceholderBody", "GroupItems", "Table.Cell", "nnn#", "UseDestinationTheme",
                   "PowerPointApi 1.9", "2 MB"):
        assert needle in t, needle


@test("docs: the duplicate-slide recipe in python-pptx.md runs and leaves no broken references")
def t_docs_duplicate():
    if not have_pptx():
        return "python-pptx not installed"
    t = open(os.path.join(ROOT, "references", "python-pptx.md"), encoding="utf-8").read()
    code = [b for b in t.split("```python")[1:] if "def duplicate" in b][0].split("```")[0]
    ns = {}
    exec(code, ns)  # the recipe text itself, so the doc cannot drift from what was tested
    from pptx import Presentation
    prs = Presentation(FIX)
    ns["duplicate"](prs, 0)
    out = tmp("dup.pptx")
    prs.save(out)
    rc, d, o = run("inspect_pptx.py", out, "--json")
    assert rc == 0 and d["totals"]["slides"] == 5 and d["slides"][4]["title"] == d["slides"][0]["title"], o
    assert not [r for r in d["risks"] if r["risk"] == "broken-slide-reference"]
    assert consistent(out)


@test("release: no em dashes in shipped files")
def t_dash():
    bad = []
    for d, _, fs in os.walk(ROOT):
        if ".git" in d or "node_modules" in d:
            continue
        for f in fs:
            if f.endswith((".md", ".py", ".json", ".txt")) or f == "LICENSE":
                if "\u2014" in open(os.path.join(d, f), encoding="utf-8").read():
                    bad.append(f)
    assert not bad, bad


@test("L5: the tracked fixture is unchanged after the whole run")
def t_fixture_untouched():
    assert sha(TRACKED) == TRACKED_SHA, "tests modified tests/fixtures/deck.pptx"


if __name__ == "__main__":
    for t in TESTS:
        t()
    shutil.rmtree(TMP, ignore_errors=True)
    c = {k: sum(1 for r in RESULTS if r[0] == k) for k in ("PASS", "FAIL", "SKIP")}
    print("\n%d passed, %d failed, %d skipped (%d tests)" % (c["PASS"], c["FAIL"], c["SKIP"], len(RESULTS)))
    sys.exit(1 if c["FAIL"] else 0)
